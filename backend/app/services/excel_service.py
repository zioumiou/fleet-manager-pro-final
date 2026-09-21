import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from typing import List, Dict, Any

def create_excel_buffer(data: List[Dict[str, Any]], sheet_name: str = "Données") -> io.BytesIO:
    """Crée un fichier Excel à partir de données"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    # Style des en-têtes
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    if data:
        # En-têtes
        headers = list(data[0].keys())
        ws.append(headers)
        
        # Style des en-têtes
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        
        # Données
        for row in data:
            ws.append([row.get(h, "") for h in headers])
        
        # Ajuster la largeur des colonnes
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output